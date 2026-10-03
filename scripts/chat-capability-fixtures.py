from pathlib import Path
from PIL import Image, ImageDraw
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas

folder = Path('docs/chat-capability-fixtures')
folder.mkdir(exist_ok=True)
im = Image.new('RGB', (600, 200), 'white')
draw = ImageDraw.Draw(im)
for index, color in enumerate(['red', 'green', 'blue']):
    draw.rectangle((index * 200, 0, (index + 1) * 200 - 1, 199), fill=color)
im.save(folder / '视觉识别测试.png')
(folder / '正文识别测试.txt').write_text('仅用于功能验收。项目暗号是：松果7319。预算金额为 24680 元。不要把测试当成真实客户资料。', encoding='utf-8')
doc = Document()
doc.add_heading('文件识别测试', 0)
doc.add_paragraph('本测试的 Word 暗号是：梧桐5826。负责人代称：测试员。')
doc.save(folder / 'Word识别测试.docx')
wb = Workbook()
ws = wb.active
ws.append(['item', 'amount'])
ws.append(['test-alpha', 130])
ws.append(['test-beta', 270])
wb.save(folder / '表格识别测试.xlsx')
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = 'PPT reading test'
slide.placeholders[1].text = 'The validation code is CEDAR-9462.'
prs.save(folder / 'PPT识别测试.pptx')
c = canvas.Canvas(str(folder / 'PDF识别测试.pdf'))
c.drawString(70, 740, 'PDF reading test. The validation code is MAPLE-8047.')
c.save()
print('Synthetic vision, TXT, Word, Excel, PPT and text-PDF fixtures created')
